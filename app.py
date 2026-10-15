import io
import re
import os
from typing import Dict, List, Tuple

import fitz
import numpy as np
import pandas as pd
import streamlit as st
from docx import Document
from sentence_transformers import SentenceTransformer

# ------------------------------------------------------------
# Resume Screening Web App
# Based on the user's main.ipynb pipeline:
# extraction -> preprocessing -> section split -> BGE embeddings
# -> section similarity -> weighted ranking -> contact extraction
# ------------------------------------------------------------

st.set_page_config(
    page_title="Resume Screening AI",
    page_icon="📄",
    layout="wide",
)

MODEL_NAME = "BAAI/bge-large-en-v1.5"
QUERY_PREFIX = "Represent this sentence for searching relevant passages: "
MAX_WORDS = 60
TOP_EVIDENCE = 3

SECTION_MAP = {
    "required_skills": ["skills", "experience", "projects"],
    "experience": ["experience", "summary"],
    "responsibilities": ["experience", "projects"],
    "preferred_skills": ["skills", "projects", "experience"],
}

SECTION_WEIGHTS = {
    "required_skills": 0.45,
    "responsibilities": 0.25,
    "experience": 0.20,
    "preferred_skills": 0.10,
}

SECTION_HEADINGS = {
    "summary": ["summary", "professional summary", "career summary", "profile",
                "about me", "objective", "career objective"],
    "skills": ["skills", "technical skills", "key skills", "core competencies",
               "technical expertise", "technologies", "tech stack", "skill set"],
    "experience": ["experience", "work experience", "professional experience",
                   "employment history", "work history", "career history", "employment"],
    "projects": ["projects", "academic projects", "personal projects",
                 "key projects", "project experience"],
    "education": ["education", "academic background", "academic qualifications",
                  "educational qualification", "qualifications"],
    "certifications": ["certifications", "certification", "courses", "training",
                       "licenses and certifications"],
    "achievements": ["achievements", "awards", "accomplishments", "honors",
                     "publications", "extracurricular"],
}

JD_HEADING_MAP = {
    "role": ["job title", "title", "role", "position", "designation",
             "about the role", "role overview", "job summary", "about the job", "overview"],
    "required_skills": ["required skills", "requirements", "must have", "must haves",
                        "required qualifications", "skills required", "key skills",
                        "technical skills", "what we are looking for", "who you are",
                        "essential skills", "mandatory skills", "technical requirements"],
    "preferred_skills": ["preferred skills", "nice to have", "good to have", "bonus",
                         "preferred qualifications", "desirable", "plus points", "added advantage"],
    "responsibilities": ["responsibilities", "key responsibilities", "what you will do",
                         "what you'll do", "duties", "job description", "your role", "day to day"],
    "experience": ["experience", "experience required", "work experience",
                   "years of experience", "eligibility"],
    "education": ["education", "qualification", "qualifications",
                  "educational qualification", "academic requirements"],
    "company": ["about us", "about the company", "who we are", "company overview"],
    "benefits": ["benefits", "what we offer", "perks", "compensation", "salary"],
}
DROP_SECTIONS = {"company", "benefits"}


@st.cache_resource(show_spinner="Loading BGE embedding model (first run may take a few minutes)...")
def load_embedding_model():
    # The same embedding model used in the notebook. No training is performed.
    return SentenceTransformer(MODEL_NAME)


def extract_uploaded_text(uploaded_file) -> str:
    name = uploaded_file.name
    ext = os.path.splitext(name)[1].lower()
    data = uploaded_file.getvalue()

    if ext == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc).strip()
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts).strip()
    if ext == ".txt":
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError(f"Unsupported file type: {ext}. Upload PDF, DOCX, or TXT.")


def preprocess(text: str) -> str:
    if not isinstance(text, str):
        return ""
    text = text.lower()
    text = text.replace("\t", " ")
    text = re.sub(r"[•‣▪◦●·∙]", "-", text)
    text = re.sub(r"http\S+|www\.\S+", " ", text)
    text = re.sub(r"[^a-z0-9+#./\-,()&@:\n ]", " ", text)
    text = re.sub(r"[ ]{2,}", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def match_resume_heading(line):
    line = line.strip().strip(":").strip("-").strip()
    if len(line.split()) > 5 or len(line) < 3:
        return None
    for section, variants in SECTION_HEADINGS.items():
        for variant in variants:
            if line == variant or line.startswith(variant):
                return section
    return None


def split_resume_sections(text: str) -> Dict[str, str]:
    sections, current, buffer = {}, "header", []
    for line in text.split("\n"):
        heading = match_resume_heading(line)
        if heading:
            if buffer:
                sections[current] = (sections.get(current, "") + "\n" + "\n".join(buffer)).strip()
            current, buffer = heading, []
        elif line.strip():
            buffer.append(line.strip())

    if buffer:
        sections[current] = (sections.get(current, "") + "\n" + "\n".join(buffer)).strip()

    sections = {k: v.strip() for k, v in sections.items() if v.strip()}
    real_sections = [k for k in sections if k != "header"]
    if len(real_sections) < 2:
        return {"full_text": text.strip()}
    return sections


def clean_heading_candidate(line: str) -> str:
    s = line.strip().strip("#").strip()
    s = s.replace("*", "").replace("_", "")
    return s.strip("-").strip("=").strip(":").lower().strip()


def lookup_jd_section(candidate: str):
    if len(candidate) < 3 or len(candidate.split()) > 6:
        return None
    for section, variants in JD_HEADING_MAP.items():
        for variant in variants:
            if candidate == variant or candidate.startswith(variant):
                return section
    return None


def match_jd_heading(line: str):
    raw = line.strip()
    if not raw or set(raw) <= set("-=_#* "):
        return None, ""
    head, sep, rest = raw.partition(":")
    if sep:
        section = lookup_jd_section(clean_heading_candidate(head))
        if section:
            return section, rest.strip()
    section = lookup_jd_section(clean_heading_candidate(raw))
    if section:
        return section, ""
    return None, ""


def split_jd_sections(text: str) -> Dict[str, str]:
    sections, current, buffer = {}, "role", []

    def flush():
        if buffer:
            sections[current] = (
                sections.get(current, "") + "\n" + "\n".join(buffer)
            ).strip()

    for line in text.split("\n"):
        section, inline = match_jd_heading(line)
        if section:
            flush()
            current = section
            buffer = [inline] if inline else []
        elif line.strip() and not set(line.strip()) <= set("-=_#* "):
            buffer.append(line.strip())

    flush()
    sections = {k: v.strip() for k, v in sections.items() if v.strip()}
    for section in DROP_SECTIONS:
        sections.pop(section, None)

    if len([k for k in sections if k != "role"]) < 1:
        # If headings are not recognized, treat the whole JD as the main query section.
        return {"required_skills": text.strip()}
    return sections


def chunk_text(text: str, max_words: int = MAX_WORDS) -> List[str]:
    chunks, buffer, count = [], [], 0
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        words = len(line.split())
        if count + words > max_words and buffer:
            chunks.append(" ".join(buffer))
            buffer, count = [], 0
        # Keep a very long line intact, matching notebook behavior.
        buffer.append(line)
        count += words
    if buffer:
        chunks.append(" ".join(buffer))
    return chunks


def embed_sections(model, sections: Dict[str, str], is_query=False):
    chunks = {section: chunk_text(text) for section, text in sections.items()}
    vectors = {}
    for section, section_chunks in chunks.items():
        if not section_chunks:
            vectors[section] = np.zeros((0, model.get_sentence_embedding_dimension()), dtype="float32")
            continue
        texts = ([QUERY_PREFIX + t for t in section_chunks] if is_query else section_chunks)
        vectors[section] = model.encode(
            texts, batch_size=16, normalize_embeddings=True,
            convert_to_numpy=True, show_progress_bar=False
        )
    return chunks, vectors


def compare_section(jd_vectors, jd_chunks, resume_vectors, resume_chunks, targets):
    if "full_text" in resume_vectors:
        targets = ["full_text"]

    matrices, texts, sources = [], [], []
    for section in targets:
        vectors = resume_vectors.get(section)
        chunks = resume_chunks.get(section, [])
        if vectors is not None and len(vectors) and chunks:
            matrices.append(vectors)
            texts.extend(chunks)
            sources.extend([section] * len(chunks))

    if not matrices or len(jd_vectors) == 0:
        return np.nan, []

    resume_matrix = np.vstack(matrices)
    if jd_vectors.shape[1] != resume_matrix.shape[1]:
        raise ValueError("Embedding dimensions differ. Check that the same embedding model is used.")

    similarities = jd_vectors @ resume_matrix.T
    best_indices = similarities.argmax(axis=1)
    best_values = similarities.max(axis=1)
    evidence = []
    for i, value in enumerate(best_values):
        j = int(best_indices[i])
        evidence.append({
            "jd_chunk": jd_chunks[i],
            "resume_chunk": texts[j],
            "resume_section": sources[j],
            "similarity": float(value),
        })
    evidence.sort(key=lambda item: item["similarity"], reverse=True)
    return float(best_values.mean()), evidence[:TOP_EVIDENCE]


def estimate_years(text: str):
    matches = re.findall(r"(?<!\d)(\d{1,2})\s*\+?\s*(?:years|year|yrs|yr)\b", text.lower())
    years = [int(v) for v in matches if 0 <= int(v) <= 60]
    return max(years) if years else None


def extract_contact(text: str):
    emails = re.findall(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", text)
    phones = []
    for candidate in re.findall(
        r"(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,5}\)?[\s\-]?)?\d{3,5}[\s\-]?\d{4,6}", text
    ):
        digits = re.sub(r"\D", "", candidate)
        if 10 <= len(digits) <= 13:
            phones.append(candidate.strip())
    linkedin = re.findall(r"linkedin\.com/in/[A-Za-z0-9\-_%]+", text, flags=re.I)
    github = re.findall(r"github\.com/[A-Za-z0-9\-_]+", text, flags=re.I)
    return {
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
        "linkedin": linkedin[0] if linkedin else "",
        "github": github[0] if github else "",
    }


def guess_name(text: str):
    for line in [line.strip() for line in text.splitlines() if line.strip()][:5]:
        if 1 < len(line.split()) <= 4 and "@" not in line and not re.search(r"\d{4}", line):
            if not re.search(r"\b(resume|curriculum vitae|linkedin|github)\b", line, re.I):
                return line
    return ""


def calculate_overall(scores):
    total, weight_total = 0.0, 0.0
    for section, weight in SECTION_WEIGHTS.items():
        value = scores.get(section, np.nan)
        if pd.notna(value):
            total += float(value) * weight
            weight_total += weight
    return total / weight_total if weight_total else np.nan


def screen_resumes(jd_text: str, uploaded_resumes, top_k: int):
    model = load_embedding_model()
    clean_jd = preprocess(jd_text)
    jd_sections = split_jd_sections(clean_jd)
    jd_chunks, jd_vectors = embed_sections(model, jd_sections, is_query=True)

    if not uploaded_resumes:
        raise ValueError("Upload at least one resume.")
    resume_records = []
    skipped = []

    for file in uploaded_resumes:
        try:
            original_text = extract_uploaded_text(file)
            if not original_text.strip():
                skipped.append(f"{file.name}: no text extracted (scanned PDF may need OCR)")
                continue
            clean_text = preprocess(original_text)
            sections = split_resume_sections(clean_text)
            chunks, vectors = embed_sections(model, sections, is_query=False)
            scores, all_evidence = {}, {}

            for jd_section, targets in SECTION_MAP.items():
                if jd_section not in jd_vectors:
                    scores[jd_section] = np.nan
                    all_evidence[jd_section] = []
                    continue
                score, evidence = compare_section(
                    jd_vectors[jd_section],
                    jd_chunks.get(jd_section, []),
                    vectors,
                    chunks,
                    targets,
                )
                scores[jd_section] = score
                all_evidence[jd_section] = evidence

            resume_records.append({
                "file_name": file.name,
                "original_text": original_text,
                "sections": sections,
                "scores": scores,
                "evidence": all_evidence,
                "raw_score": calculate_overall(scores),
                "years": estimate_years(original_text),
                **extract_contact(original_text),
                "name": guess_name(original_text),
            })
        except Exception as exc:
            skipped.append(f"{file.name}: {exc}")

    if not resume_records:
        raise ValueError("No readable resumes could be screened. " + " | ".join(skipped))

    results = pd.DataFrame(resume_records)
    results = results[results["raw_score"].notna()].copy()
    if results.empty:
        raise ValueError(
            "No valid similarity scores were produced. The JD may not have detectable sections, "
            "or the resumes may not contain extractable text."
        )

    low, high = results["raw_score"].min(), results["raw_score"].max()
    if high - low < 1e-6:
        results["relative_score"] = 50.0
    else:
        results["relative_score"] = ((results["raw_score"] - low) / (high - low) * 100).round(1)

    results = results.sort_values("raw_score", ascending=False).reset_index(drop=True)
    results.insert(0, "rank", np.arange(1, len(results) + 1))
    return results.head(top_k), results, jd_sections, skipped


# ------------------------------ UI ------------------------------

st.title("📄 Resume Screening AI")
st.caption("Upload a job description and resumes to compare them using the BGE embeddings from your notebook.")

with st.sidebar:
    st.header("Screening settings")
    top_k = st.slider("Candidates to shortlist", min_value=1, max_value=50, value=5)
    st.markdown("---")
    st.markdown("**Embedding model**")
    st.code(MODEL_NAME, language="text")
    st.caption("First run downloads the model. The app does not retrain it.")

left, right = st.columns([1, 1])
with left:
    st.subheader("1. Job description")
    jd_file = st.file_uploader("Upload JD (TXT, PDF, DOCX)", type=["txt", "pdf", "docx"], key="jd")
    jd_text_input = st.text_area("Or paste the job description", height=220, placeholder="Paste the full job description here...")

with right:
    st.subheader("2. Resumes")
    resume_files = st.file_uploader(
        "Upload multiple resumes (PDF, DOCX, TXT)",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        key="resumes",
    )
    st.info("Tip: Select several files at once. Text-based PDFs work best; scanned PDFs may require OCR.")

if st.button("🚀 Screen resumes", type="primary", use_container_width=True):
    if jd_text_input.strip():
        final_jd = jd_text_input.strip()
    elif jd_file is not None:
        try:
            final_jd = extract_uploaded_text(jd_file)
        except Exception as exc:
            st.error(f"Could not read the job description: {exc}")
            st.stop()
    else:
        st.error("Upload a job description or paste its text.")
        st.stop()

    if not resume_files:
        st.error("Upload at least one resume.")
        st.stop()

    try:
        with st.spinner("Extracting resumes, embedding text, and calculating similarities..."):
            shortlist, all_results, jd_sections, skipped = screen_resumes(final_jd, resume_files, top_k)

        st.success(f"Screening completed. Ranked {len(all_results)} readable resumes.")
        if skipped:
            with st.expander(f"Files skipped ({len(skipped)})"):
                for item in skipped:
                    st.write("• " + item)

        st.subheader("3. Ranked shortlist")
        st.caption(
            "Relative score scales candidates against the resumes in this run. "
            "It is not an absolute probability of job fit and can change when the candidate pool changes."
        )
        display_columns = ["rank", "name", "file_name", "relative_score", "raw_score", "years",
                           "email", "phone", "linkedin", "github"]
        table = shortlist[display_columns].rename(columns={
            "relative_score": "Relative score / 100",
            "raw_score": "Raw similarity",
            "file_name": "Resume file",
        })
        st.dataframe(table, use_container_width=True, hide_index=True)

        csv_data = all_results[display_columns].to_csv(index=False).encode("utf-8")
        st.download_button(
            "⬇️ Download all results (CSV)",
            data=csv_data,
            file_name="resume_screening_results.csv",
            mime="text/csv",
        )

        st.subheader("4. Candidate details and evidence")
        for _, candidate in shortlist.iterrows():
            label = f"#{int(candidate['rank'])} · {candidate['name'] or candidate['file_name']} · {candidate['relative_score']:.1f}/100 relative"
            with st.expander(label):
                a, b, c = st.columns(3)
                a.metric("Relative score", f"{candidate['relative_score']:.1f}/100")
                b.metric("Raw similarity", f"{candidate['raw_score']:.4f}")
                c.metric("Estimated years", str(candidate["years"]) if pd.notna(candidate["years"]) else "Not found")

                st.write("**Contact details**")
                st.write({
                    "Email": candidate["email"] or "Not found",
                    "Phone": candidate["phone"] or "Not found",
                    "LinkedIn": candidate["linkedin"] or "Not found",
                    "GitHub": candidate["github"] or "Not found",
                })

                for section, items in candidate["evidence"].items():
                    if not items:
                        continue
                    st.markdown(f"**{section.replace('_', ' ').title()} evidence**")
                    for item in items:
                        st.markdown(f"- Similarity `{item['similarity']:.3f}` · resume section: `{item['resume_section']}`")
                        st.write("JD:", item["jd_chunk"])
                        st.write("Resume:", item["resume_chunk"])
                        st.divider()

        with st.expander("Detected job-description sections"):
            st.json(jd_sections)

    except Exception as exc:
        st.error(f"Screening failed: {exc}")
        st.exception(exc)

st.markdown("---")
st.caption("Review candidates manually. Similarity scores and extracted contact details can be imperfect; do not use them as the sole basis for hiring decisions.")
